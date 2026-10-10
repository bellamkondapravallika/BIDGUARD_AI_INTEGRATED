import csv
import io
import json
import logging
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

logger = logging.getLogger(__name__)

# Patterns for GeM compliance and statutory verifications
PAN_PATTERN = re.compile(r"\b[A-Z]{5}[0-9]{4}[A-Z]\b")
GST_PATTERN = re.compile(r"\b\d{2}[A-Z]{5}\d{4}[A-Z]{1}[A-Z0-9]{1}[Z]{1}[A-Z0-9]{1}\b")
UDYAM_PATTERN = re.compile(r"\bUDYAM-[A-Z]{2}-\d{2}-\d{7}\b", re.IGNORECASE)
CIN_PATTERN = re.compile(r"\b[UL]\d{5}[A-Z]{2}\d{4}[A-Z]{3}\d{6}\b")
DATE_PATTERN = re.compile(r"\b(?:\d{1,2}[-/\.]\d{1,2}[-/\.]\d{2,4}|\d{4}[-/\.]\d{1,2}[-/\.]\d{1,2})\b")
TURNOVER_PATTERN = re.compile(r"(?:turnover|revenue|annual\s+sales)[^\d\n\r]{1,25}(?:₹|rs\.?|inr)?\s*([\d,]+(?:\.\d+)?)\s*(cr|crore|crores|lakh|lakhs|k)?", re.IGNORECASE)
EXP_PATTERN = re.compile(r"(\d+)\s*(?:\+|\s*plus)?\s*(?:years|yrs)\s*(?:of)?\s*(?:experience|standing|proven)", re.IGNORECASE)


class DocumentExtractor:

    @staticmethod
    def extract(file_path: str, file_type: str) -> Dict[str, Any]:
        """
        Extract text, pages, and structured procurement metadata from a supported document.
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        file_type = file_type.lower().strip(".")
        full_text = ""
        pages: List[Dict[str, Any]] = []

        try:
            if file_type == "pdf":
                full_text, pages = DocumentExtractor._extract_pdf(path)
            elif file_type in ["docx", "doc"]:
                full_text, pages = DocumentExtractor._extract_docx(path)
            elif file_type in ["xlsx", "xls"]:
                full_text, pages = DocumentExtractor._extract_xlsx(path)
            elif file_type == "csv":
                full_text, pages = DocumentExtractor._extract_csv(path)
            elif file_type in ["txt", "text"]:
                full_text, pages = DocumentExtractor._extract_txt(path)
            else:
                full_text, pages = DocumentExtractor._extract_txt(path)
        except Exception as exc:
            logger.error("Error extracting text from %s: %s", path, exc)
            return {
                "success": False,
                "error": str(exc),
                "full_text": "",
                "pages": [],
                "entities": {},
            }

        entities = DocumentExtractor._extract_entities(full_text)

        return {
            "success": True,
            "error": None,
            "full_text": full_text,
            "pages": pages,
            "page_count": len(pages),
            "entities": entities,
        }

    @staticmethod
    def _extract_pdf(path: Path) -> tuple[str, List[Dict[str, Any]]]:
        import pypdf

        full_text = []
        pages = []

        with open(path, "rb") as f:
            reader = pypdf.PdfReader(f)
            for idx, page in enumerate(reader.pages, start=1):
                page_text = page.extract_text() or ""
                clean_text = page_text.strip()
                pages.append({"page": idx, "text": clean_text})
                if clean_text:
                    full_text.append(clean_text)

        combined = "\n\n".join(full_text)
        return combined, pages

    @staticmethod
    def _extract_docx(path: Path) -> tuple[str, List[Dict[str, Any]]]:
        import docx

        doc = docx.Document(path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        combined = "\n".join(paragraphs)
        pages = [{"page": 1, "text": combined}]
        return combined, pages

    @staticmethod
    def _extract_xlsx(path: Path) -> tuple[str, List[Dict[str, Any]]]:
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            lines.append(f"--- Sheet: {sheet_name} ---")
            for row in sheet.iter_rows(values_only=True):
                non_empty = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if non_empty:
                    lines.append(" | ".join(non_empty))

        combined = "\n".join(lines)
        pages = [{"page": 1, "text": combined}]
        return combined, pages

    @staticmethod
    def _extract_csv(path: Path) -> tuple[str, List[Dict[str, Any]]]:
        lines = []
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            for row in reader:
                if any(cell.strip() for cell in row):
                    lines.append(" | ".join(cell.strip() for cell in row))

        combined = "\n".join(lines)
        pages = [{"page": 1, "text": combined}]
        return combined, pages

    @staticmethod
    def _extract_txt(path: Path) -> tuple[str, List[Dict[str, Any]]]:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read().strip()
        pages = [{"page": 1, "text": text}]
        return text, pages

    @staticmethod
    def _extract_entities(text: str) -> Dict[str, Any]:
        """
        Identify high-confidence procurement compliance entities.
        """
        pans = list(set(PAN_PATTERN.findall(text)))
        gsts = list(set(GST_PATTERN.findall(text)))
        udyams = list(set(UDYAM_PATTERN.findall(text)))
        cins = list(set(CIN_PATTERN.findall(text)))
        dates = list(set(DATE_PATTERN.findall(text)))

        # Turnover
        turnover_vals: List[float] = []
        for match in TURNOVER_PATTERN.finditer(text):
            num_str, unit = match.groups()
            try:
                val = float(num_str.replace(",", ""))
                if unit:
                    unit_lower = unit.lower()
                    if "cr" in unit_lower:
                        val *= 10_000_000
                    elif "lakh" in unit_lower:
                        val *= 100_000
                    elif "k" in unit_lower:
                        val *= 1_000
                turnover_vals.append(val)
            except Exception:
                pass

        # Experience
        experience_vals: List[int] = []
        for match in EXP_PATTERN.finditer(text):
            try:
                experience_vals.append(int(match.group(1)))
            except Exception:
                pass

        # OEM Authorization indicator
        has_oem_auth = bool(re.search(r"\b(?:oem\s+authorization|authorized\s+partner|channel\s+partner|manufacturer\s+authorization|maf)\b", text, re.IGNORECASE))

        # Blacklist/debarment check
        has_debarment_declaration = bool(re.search(r"\b(?:not\s+blacklisted|never\s+been\s+blacklisted|no\s+debarment|debarment\s+declaration)\b", text, re.IGNORECASE))

        return {
            "pan_numbers": pans,
            "gstin_numbers": gsts,
            "udyam_numbers": udyams,
            "cin_numbers": cins,
            "dates": dates[:10],
            "turnover_values": turnover_vals,
            "experience_years": experience_vals,
            "has_oem_authorization": has_oem_auth,
            "has_debarment_declaration": has_debarment_declaration,
        }
